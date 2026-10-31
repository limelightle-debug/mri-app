"""The natural rate of interest (r*) for the institute's paper: data/rstar.json.

The New York Fed's estimates, quarterly, one-sided (as known at each date), US:
- HLW: Holston, Laubach and Williams ("Measuring the Natural Rate of Interest after COVID-19", 2023)
- LW: Laubach and Williams (2003), the baseline model
with HLW's inputs for the federal funds rate (quarterly average, annualized) and the expected inflation it uses,
so the real policy rate is ff - pie. Updated by the New York Fed four times a year; read at most once a day.
{"hist": [["YYYY-Qn", hlw, lw, ff, pie], ...], "final": "YYYY-Qn", "asof": note, "updatedAt", "checkedAt"}
"""
import io
import json
import sys
import time
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
import requests

OUT = Path(__file__).resolve().parent.parent / "data" / "rstar.json"
BASE = "https://www.newyorkfed.org/medialibrary/media/research/economists/williams/data/"
UA = ("Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/124.0 Safari/537.36")


def book(name):
    r = requests.get(BASE + name, headers={"User-Agent": UA}, timeout=60)
    r.raise_for_status()
    return openpyxl.load_workbook(io.BytesIO(r.content), data_only=True)


def q(d):
    return f"{d.year}-Q{(d.month - 1) // 3 + 1}"


def num(v):
    return round(float(v), 3) if isinstance(v, (int, float)) else None


def main(force=False):
    old = json.loads(OUT.read_text()) if OUT.exists() else {}
    if not force and old.get("checkedAt") and time.time() - datetime.fromisoformat(
            old["checkedAt"].replace("Z", "+00:00")).timestamp() < 86400:
        return
    hlw, inp, lw, note = {}, {}, {}, ""
    wb = book("Holston_Laubach_Williams_current_estimates.xlsx")
    for row in wb["HLW Estimates"].iter_rows(values_only=True):
        if row and isinstance(row[0], str) and row[0].startswith("Final data point"):
            note = row[0]
        if row and isinstance(row[0], datetime):
            hlw[q(row[0])] = num(row[10])  # Natural Rate (r*), US
    for row in wb["US input data"].iter_rows(values_only=True):
        if row and isinstance(row[0], datetime):
            inp[q(row[0])] = (num(row[4]), num(row[3]))  # interest, inflation.expectations
    wl = book("Laubach_Williams_current_estimates.xlsx")
    for row in wl["data"].iter_rows(values_only=True):
        if row and isinstance(row[0], str) and row[0].startswith("Estimates based on"):
            note = row[0]
        if row and isinstance(row[0], datetime):
            lw[q(row[0])] = num(row[2])  # one-sided rstar
    keys = sorted(set(hlw) | set(lw))
    hist = [[k, hlw.get(k), lw.get(k), *(inp.get(k) or (None, None))] for k in keys]
    if len(hist) < 100:
        sys.exit(f"rstar: only {len(hist)} quarters")
    now = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    same = old.get("hist") == hist
    OUT.write_text(json.dumps({"hist": hist, "final": hist[-1][0], "asof": note,
                               "source": "Federal Reserve Bank of New York (HLW, LW; one-sided estimates)",
                               "updatedAt": old.get("updatedAt", now) if same else now, "checkedAt": now},
                              separators=(",", ":")))
    if not same:
        print(hist[-1])


if __name__ == "__main__":
    main("--force" in sys.argv)
